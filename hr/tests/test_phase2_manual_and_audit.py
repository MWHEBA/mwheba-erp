"""
اختبارات شاملة للمرحلة 2:
1. الإدخال والتعديل اليدوي للحضور (Manual Attendance)
2. سجل تدقيق التعديلات اليدوية (AttendanceAuditLog)
3. حماية الفترات المقفلة والرواتب المعتمدة
4. استيراد وتوليد قالب Excel المجمع مع التقرير التفصيلي
5. تسوية قطع الإجازات (Leave Cutoff Reconciliation)
6. الصلاحيات المخصصة RBAC
"""
import io
import json
import pytest
from datetime import date, time, datetime, timedelta
from decimal import Decimal
import openpyxl

from django.test import RequestFactory
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType

from hr.models import (
    Employee, Department, JobTitle, Shift, Attendance,
    AttendanceAuditLog, LeaveType, Leave, LeaveSummary,
    Payroll, Contract
)
from hr.views.attendance_views import (
    attendance_manual_save,
    attendance_import_template,
    attendance_import_excel,
    attendance_audit_history,
)

User = get_user_model()


@pytest.fixture
def phase2_setup(db):
    """إعداد بيانات الاختبار الأساسية للمرحلة 2"""
    dept = Department.objects.create(name_ar='إدارة تكنولوجيا المعلومات', code='IT')
    job = JobTitle.objects.create(title_ar='مهندس برمجيات', department=dept)
    shift = Shift.objects.create(
        name='وردية الصباح',
        start_time=time(9, 0),
        end_time=time(17, 0)
    )

    user = User.objects.create_user(username='hr_admin_phase2', password='password123', is_staff=True)
    
    # منح الصلاحيات المخصصة
    content_type = ContentType.objects.get_for_model(Attendance)
    perm_manual, _ = Permission.objects.get_or_create(
        codename='can_manual_attendance',
        content_type=content_type,
        defaults={'name': 'يمكنه إضافة وتعديل الحضور يدوياً'}
    )
    perm_import, _ = Permission.objects.get_or_create(
        codename='can_import_attendance',
        content_type=content_type,
        defaults={'name': 'يمكنه استيراد الحضور من Excel'}
    )
    user.user_permissions.add(perm_manual, perm_import)

    emp = Employee.objects.create(
        name='طارق محمود علام',
        employee_number='EMP-P2-001',
        national_id='29501011234567',
        birth_date=date(1995, 1, 1),
        gender='male',
        marital_status='single',
        mobile_phone='01012345678',
        created_by=user,
        department=dept,
        job_title=job,
        shift=shift,
        status='active',
        hire_date=date(2025, 1, 1),
    )

    supervisor = Employee.objects.create(
        name='مدير الموقع المشرف',
        employee_number='EMP-SUP-001',
        national_id='29001011234568',
        birth_date=date(1990, 1, 1),
        gender='male',
        marital_status='married',
        mobile_phone='01098765432',
        created_by=user,
        department=dept,
        job_title=job,
        status='active',
        hire_date=date(2024, 1, 1),
    )

    Contract.objects.create(
        employee=emp,
        contract_number='CNT-P2-001',
        basic_salary=Decimal('6000.00'),
        start_date=date(2025, 1, 1),
        status='active',
        created_by=user
    )

    return {
        'dept': dept,
        'job': job,
        'shift': shift,
        'user': user,
        'emp': emp,
        'supervisor': supervisor,
        'rf': RequestFactory()
    }


@pytest.mark.django_db
class TestPhase2ManualAttendanceAndAudit:
    """اختبارات الإدخال والتعديل اليدوي وسجل التدقيق"""

    def test_manual_attendance_create_and_audit_log(self, phase2_setup):
        """اختبار إنشاء حضور يدوي وتوثيقه في سجل التدقيق"""
        rf = phase2_setup['rf']
        user = phase2_setup['user']
        emp = phase2_setup['emp']
        supervisor = phase2_setup['supervisor']

        target_date = date(2026, 5, 10)
        request = rf.post('/hr/attendance/manual-save/', {
            'employee_id': emp.id,
            'date': target_date.strftime('%Y-%m-%d'),
            'status': 'present',
            'check_in': '09:00',
            'check_out': '17:00',
            'supervisor_id': supervisor.id,
            'reason': 'عطل طارئ في ماكينة البصمة الرئيسية',
            'notes': 'تم التحقق من المشرف',
        })
        request.user = user

        response = attendance_manual_save(request)
        assert response.status_code == 200
        data = json.loads(response.content.decode('utf-8'))
        assert data['success'] is True

        att = Attendance.objects.get(employee=emp, date=target_date)
        assert att.status == 'present'
        assert att.is_manual is True
        assert att.source == 'manual'
        assert att.manual_created_by == user
        assert att.work_hours == Decimal('8.00')

        # التحقق من سجل التدقيق
        audit = AttendanceAuditLog.objects.filter(attendance=att).first()
        assert audit is not None
        assert audit.field_name == 'إنشاء يدوي'
        assert audit.changed_by == user
        assert audit.supervisor_witness == supervisor
        assert 'عطل طارئ' in audit.reason

    def test_manual_attendance_edit_captures_diff(self, phase2_setup):
        """اختبار تعديل حضور موجود وتوثيق الفوارق في سجل التدقيق"""
        rf = phase2_setup['rf']
        user = phase2_setup['user']
        emp = phase2_setup['emp']

        target_date = date(2026, 5, 11)
        att = Attendance.objects.create(
            employee=emp,
            date=target_date,
            shift=phase2_setup['shift'],
            status='absent',
            work_hours=Decimal('0')
        )

        request = rf.post('/hr/attendance/manual-save/', {
            'employee_id': emp.id,
            'date': target_date.strftime('%Y-%m-%d'),
            'status': 'late',
            'check_in': '09:30',
            'check_out': '17:00',
            'reason': 'تعديل غياب إلى حضور متأخر بناء على تصريح المدير',
        })
        request.user = user

        response = attendance_manual_save(request)
        assert response.status_code == 200

        att.refresh_from_db()
        assert att.status == 'late'
        assert att.work_hours == Decimal('7.50')
        assert att.late_minutes == Decimal('30.00')

        # التحقق من تسجيل الحقول المعدلة في التدقيق
        audit_records = AttendanceAuditLog.objects.filter(attendance=att)
        assert audit_records.count() >= 1
        status_audit = audit_records.filter(field_name='الحالة').first()
        assert status_audit is not None
        assert status_audit.old_value == 'absent'
        assert status_audit.new_value == 'late'

    def test_closed_payroll_blocks_direct_manual_edit(self, phase2_setup):
        """اختبار حظر التعديل المباشر لأي شهر مغلق/معتمد لحماية القيود المحاسبية"""
        rf = phase2_setup['rf']
        user = phase2_setup['user']
        emp = phase2_setup['emp']

        target_date = date(2026, 3, 15)
        payroll_month = date(2026, 3, 1)

        contract = emp.contracts.first()
        Payroll.objects.create(
            employee=emp,
            contract=contract,
            month=payroll_month,
            status='approved',
            basic_salary=Decimal('6000.00'),
            gross_salary=Decimal('6000.00'),
            net_salary=Decimal('6000.00'),
            processed_by=user
        )

        request = rf.post('/hr/attendance/manual-save/', {
            'employee_id': emp.id,
            'date': target_date.strftime('%Y-%m-%d'),
            'status': 'present',
            'check_in': '09:00',
            'check_out': '17:00',
            'reason': 'محاولة تعديل شهر مغلق',
        })
        request.user = user

        response = attendance_manual_save(request)
        assert response.status_code == 400
        data = json.loads(response.content.decode('utf-8'))
        assert data['success'] is False
        assert 'راتب هذا الشهر معتمد' in data['message']


@pytest.mark.django_db
class TestPhase2ExcelImport:
    """اختبارات استيراد Excel بنظام التقرير التفصيلي"""

    def test_attendance_import_template_generation(self, phase2_setup):
        """اختبار توليد وتحميل قالب Excel النموذجي"""
        rf = phase2_setup['rf']
        user = phase2_setup['user']

        request = rf.get('/hr/attendance/import/template/')
        request.user = user

        response = attendance_import_template(request)
        assert response.status_code == 200
        assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'attendance_import_template.xlsx' in response['Content-Disposition']

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert "كود الموظف" in headers
        assert "التاريخ (YYYY-MM-DD)" in headers

    def test_attendance_import_excel_with_valid_and_invalid_rows(self, phase2_setup):
        """اختبار استيراد Excel مجمع بصفوف صالحة وخاطئة واستخراج التقرير التفصيلي"""
        rf = phase2_setup['rf']
        user = phase2_setup['user']
        emp = phase2_setup['emp']

        # إنشاء ملف Excel في الذاكرة
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["كود الموظف", "اسم الموظف", "التاريخ", "وقت الحضور", "وقت الانصراف", "ملاحظات"])
        # صف صحيح 1
        ws.append([emp.employee_number, emp.name, "2026-05-18", "09:00", "17:00", "صف سليم"])
        # صف صحيح 2
        ws.append([emp.employee_number, emp.name, "2026-05-19", "09:15", "17:15", "صف سليم 2"])
        # صف خاطئ: كود موظف غير موجود
        ws.append(["INVALID-EMP-999", "موظف مجهول", "2026-05-20", "09:00", "17:00", "خطأ كود"])
        # صف خاطئ: صيغة تاريخ غير صحيحة
        ws.append([emp.employee_number, emp.name, "INVALID-DATE", "09:00", "17:00", "خطأ تاريخ"])

        excel_io = io.BytesIO()
        wb.save(excel_io)
        excel_io.seek(0)

        from django.core.files.uploadedfile import SimpleUploadedFile
        upload_file = SimpleUploadedFile("test_attendance.xlsx", excel_io.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        request = rf.post('/hr/attendance/import/excel/', {'excel_file': upload_file})
        request.user = user

        response = attendance_import_excel(request)
        assert response.status_code == 200
        data = json.loads(response.content.decode('utf-8'))
        assert data['success'] is True
        assert data['total_rows'] == 4
        assert data['success_count'] == 2
        assert data['error_count'] == 2

        # التحقق من تفاصيل الأخطاء
        errors = data['errors']
        assert len(errors) == 2
        assert errors[0]['row'] == 4
        assert 'INVALID-EMP-999' in errors[0]['code']
        assert errors[1]['row'] == 5
        assert 'صيغة التاريخ' in errors[1]['reason']

        # التأكد من حفظ الصفوف السليمة
        att1 = Attendance.objects.get(employee=emp, date=date(2026, 5, 18))
        assert att1.status == 'present'
        assert att1.source == 'excel'
        assert att1.work_hours == Decimal('8.00')


@pytest.mark.django_db
class TestPhase2LeaveCutoffReconciliation:
    """اختبارات تسوية قطع الإجازة عند حضور الموظف الفعلي"""

    def test_leave_cutoff_reconciliation_in_leave_summary(self, phase2_setup):
        """إذا نزل الموظف للعمل أثناء إجازته، يتم استثناء الأيام التي حضرها من الخصم والعد"""
        emp = phase2_setup['emp']
        
        user = phase2_setup['user']
        shift = phase2_setup['shift']
        leave_type = LeaveType.objects.create(
            name_ar='إجازة سنوية',
            code='ANNUAL_P2',
            category='annual',
            max_days_per_year=21,
            is_paid=True,
            deduction_multiplier=Decimal('1.0')
        )

        # إجازة 5 أيام من 10 إلى 14 مايو 2026
        Leave.objects.create(
            employee=emp,
            leave_type=leave_type,
            start_date=date(2026, 5, 10),
            end_date=date(2026, 5, 14),
            days_count=5,
            status='approved',
            approved_by=user
        )

        # الموظف حضر فعلياً يومين من أيام الإجازة (13 و 14 مايو)
        Attendance.objects.create(
            employee=emp,
            date=date(2026, 5, 13),
            shift=shift,
            status='present',
            work_hours=Decimal('8.0')
        )
        Attendance.objects.create(
            employee=emp,
            date=date(2026, 5, 14),
            shift=shift,
            status='late',
            work_hours=Decimal('7.5')
        )

        # حساب ملخص الإجازات
        summary = LeaveSummary.objects.create(
            employee=emp,
            month=date(2026, 5, 1)
        )
        summary.calculate()

        # يجب أن تُحسب الإجازة 3 أيام فقط (10، 11، 12) وليس 5 أيام
        assert summary.annual_leave_days == 3
        assert summary.total_paid_days == 3

